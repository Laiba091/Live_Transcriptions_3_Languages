"""Saves sessions as Excel in data/sessions/excel/ (Transcript + Summary sheets)."""
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill

from app.core.config import settings

from .schema import EXCEL_COLUMNS

RTL_LANGUAGES = {"ur", "ps"}
COLUMN_WIDTHS = {"id": 6, "clock_time": 20, "language": 9, "language_name": 12, "confidence": 11,
                 "text": 70, "normalized_text": 60, "model": 28}


class ExcelWriter:
    def __init__(self, folder: Path | None = None):
        self.folder = folder or settings.excel_dir
        self.folder.mkdir(parents=True, exist_ok=True)

    def path_for(self, session_id: str) -> Path:
        return self.folder / f"session_{session_id}.xlsx"

    def save(self, session) -> Path:
        path = self.path_for(session.id)
        transcript = pd.DataFrame([r.to_dict() for r in session.records], columns=EXCEL_COLUMNS)
        summary = pd.DataFrame([{"language": k, **v} for k, v in session.summary().items()],
                               columns=["language", "segments", "seconds"])

        with pd.ExcelWriter(path, engine="openpyxl") as writer:
            transcript.to_excel(writer, sheet_name="Transcript", index=False)
            summary.to_excel(writer, sheet_name="Summary", index=False)
            self._format(writer.sheets["Transcript"])
            self._format(writer.sheets["Summary"])
        return path

    @staticmethod
    def _format(ws) -> None:
        header_fill = PatternFill("solid", fgColor="1F2937")
        headers = {}
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
            headers[cell.value] = cell.column_letter
            ws.column_dimensions[cell.column_letter].width = COLUMN_WIDTHS.get(cell.value, 12)
        ws.freeze_panes = "A2"

        if "language" not in headers:
            return
        lang_col = headers["language"]
        for col_name in ("text", "normalized_text"):
            col = headers.get(col_name)
            if not col:
                continue
            for row in range(2, ws.max_row + 1):
                rtl = ws[f"{lang_col}{row}"].value in RTL_LANGUAGES
                ws[f"{col}{row}"].alignment = Alignment(
                    wrap_text=True, horizontal="right" if rtl else "left", readingOrder=2 if rtl else 1)
